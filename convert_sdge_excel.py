"""
SDGE AMI 2.0 Excel to JSON Conversion Script

Reads: SDGE AMI 2.0 Program - Business Process Inventory _ Visualization.xlsx
Sheet: "Inventory List " (note trailing space)
Columns: ID, L1, L2, L3, L4, Fuel Coverage

Outputs:
- hierarchy-data.json: Nested structure L1 → L2 → L3 → L4 (leaf nodes include id, fuel_coverage)
- search-index.json: Flat array for search across process names
"""

import json
import os
from collections import OrderedDict

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl...")
    os.system("pip install openpyxl")
    import openpyxl

EXCEL_FILE = "SDGE AMI 2.0 Program - Business Process Inventory _ Visualization.xlsx"
SHEET_NAME = "Inventory List "  # Note trailing space
COLUMNS = ["ID", "L1", "L2", "L3", "L4", "Fuel Coverage"]


def find_sheet(workbook):
    """Find sheet by exact name or by name without trailing space."""
    for name in workbook.sheetnames:
        if name == SHEET_NAME or name.strip() == "Inventory List":
            return workbook[name]
    # Fallback to first sheet
    return workbook.active


def get_column_indices(header_row):
    """Map column names to 1-based indices."""
    indices = {}
    for idx, cell in enumerate(header_row, start=1):
        if cell and cell.value:
            val = str(cell.value).strip()
            for col in COLUMNS:
                if val == col or val.lower().replace(" ", "") == col.lower().replace(" ", ""):
                    indices[col] = idx
                    break
    return indices


def build_hierarchy(rows):
    """Build nested L1 → L2 → L3 → L4 structure."""
    root = {"name": "Process Hierarchy", "children": []}
    # Use OrderedDict to preserve insertion order for consistent nesting
    l1_map = OrderedDict()

    for row in rows:
        row_id = row.get("ID")
        l1_name = (row.get("L1") or "").strip()
        l2_name = (row.get("L2") or "").strip()
        l3_name = (row.get("L3") or "").strip()
        l4_name = (row.get("L4") or "").strip()
        fuel = (row.get("Fuel Coverage") or "").strip() or "Both"

        if not l1_name:
            continue

        # L1
        if l1_name not in l1_map:
            l1_map[l1_name] = {"name": l1_name, "level": "L1", "children": []}
            root["children"].append(l1_map[l1_name])

        l1_node = l1_map[l1_name]
        l2_map = {c["name"]: c for c in l1_node["children"] if "children" in c}

        # L2
        if l2_name and l2_name not in l2_map:
            l2_node = {"name": l2_name, "level": "L2", "children": []}
            l1_node["children"].append(l2_node)
            l2_map[l2_name] = l2_node
        elif l2_name:
            l2_node = l2_map[l2_name]
        else:
            continue

        l3_map = {c["name"]: c for c in l2_node["children"] if "children" in c}

        # L3
        if l3_name and l3_name not in l3_map:
            l3_node = {"name": l3_name, "level": "L3", "children": []}
            l2_node["children"].append(l3_node)
            l3_map[l3_name] = l3_node
        elif l3_name:
            l3_node = l3_map[l3_name]
        else:
            continue

        # L4 (leaf)
        if l4_name:
            try:
                id_val = int(float(row_id)) if row_id is not None and row_id != "" else None
            except (ValueError, TypeError):
                id_val = row_id
            l4_node = {
                "name": l4_name,
                "level": "L4",
                "id": id_val,
                "fuel_coverage": fuel if fuel in ("Both", "Electric") else "Both",
            }
            l3_node["children"].append(l4_node)

    return root


def build_search_index(root):
    """Build flat search index from hierarchy."""
    index = []

    def traverse(node, path=None):
        path = path or []
        if node.get("level"):
            full_path = " > ".join(path + [node["name"]])
            entry = {
                "name": node["name"],
                "level": node["level"],
                "path": full_path,
                "parent": path[-1] if path else "",
            }
            if node["level"] == "L4":
                entry["id"] = node.get("id")
                entry["fuel_coverage"] = node.get("fuel_coverage", "Both")
            index.append(entry)

        for child in node.get("children", []):
            traverse(child, path + [node["name"]] if node.get("name") else path)

    for child in root.get("children", []):
        traverse(child)

    return index


def main():
    if not os.path.exists(EXCEL_FILE):
        print(f"Error: File not found: {EXCEL_FILE}")
        print("Please place the Excel file in the same folder as this script.")
        return 1

    wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    sheet = find_sheet(wb)

    # Read header and data
    rows_data = []
    col_map = {}
    all_rows = list(sheet.iter_rows(min_row=1, values_only=True))

    # Build column map from first row
    if all_rows:
        header = all_rows[0]
        for i, v in enumerate(header):
            if v is not None and str(v).strip():
                name = str(v).strip()
                if name == "ID" or (len(name) <= 4 and "ID" in name.upper()):
                    col_map["ID"] = i
                elif name == "L1":
                    col_map["L1"] = i
                elif name == "L2":
                    col_map["L2"] = i
                elif name == "L3":
                    col_map["L3"] = i
                elif name == "L4":
                    col_map["L4"] = i
                elif "Fuel" in name or "fuel" in name.lower():
                    col_map["Fuel Coverage"] = i
    if not col_map:
        col_map = {"ID": 0, "L1": 1, "L2": 2, "L3": 3, "L4": 4, "Fuel Coverage": 5}

    for row in all_rows[1:]:
        row_dict = {}
        for k, idx in col_map.items():
            if idx < len(row):
                val = row[idx]
                row_dict[k] = val if val is None else str(val).strip()
        if any(row_dict.get(c) for c in ["L1", "L2", "L3", "L4"]):
            rows_data.append(row_dict)

    wb.close()

    hierarchy = build_hierarchy(rows_data)
    search_index = build_search_index(hierarchy)

    with open("hierarchy-data.json", "w", encoding="utf-8") as f:
        json.dump(hierarchy, f, indent=2, ensure_ascii=False)

    with open("search-index.json", "w", encoding="utf-8") as f:
        json.dump(search_index, f, indent=2, ensure_ascii=False)

    print("Conversion complete.")
    print(f"  - hierarchy-data.json: {len(hierarchy.get('children', []))} L1 processes")
    print(f"  - search-index.json: {len(search_index)} entries")
    return 0


if __name__ == "__main__":
    exit(main())
